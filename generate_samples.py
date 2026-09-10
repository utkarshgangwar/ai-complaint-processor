from pathlib import Path
import docx
from fpdf import FPDF

# Ensure data directory exists
DATA_DIR = Path("data")
DATA_DIR.mkdir(parents=True, exist_ok=True)

# -------------------------------------------------------------
# Document 1: Text File (.txt) - Billing & Escalation Issue
# -------------------------------------------------------------
doc1_content = """CUSTOMER INCIDENT & COMPLAINT FORM
Case Reference: INC-88219
Date: 2026-08-14
Customer Name: Rajesh Sharma
Email: rajesh.sharma@example.com
Phone: +91-9876543210
Account ID: ACCT-90812

Incident Details:
I am writing to formally report an unauthorized recurring charge of Rs. 4,999 debited from my account on August 12, 2026, under invoice reference #INV-88219. I canceled this cloud storage subscription back in June 2026 and have the email confirmation receipt saved.

Despite speaking to frontline phone support twice yesterday, no refund transaction reference was issued. If this amount is not reversed within 48 business hours, I will be forced to file a dispute with my banking ombudsman and seek consumer legal redressal.

Action Taken So Far:
Frontline support advised that the billing desk is reviewing the cancellation log. No refund processed yet. Case remains unresolved.
"""

file1_path = DATA_DIR / "complaint_001.txt"
with open(file1_path, "w", encoding="utf-8") as f:
    f.write(doc1_content)
print(f"Generated: {file1_path}")


# -------------------------------------------------------------
# Document 2: Word Document (.docx) - Critical Service Outage
# -------------------------------------------------------------
doc2 = docx.Document()
doc2.add_heading("SERVICE INTERRUPTION & SLA BREACH REPORT", level=1)

p1 = doc2.add_paragraph()
p1.add_run("Customer Name: ").bold = True
p1.add_run("Priya Nair\n")
p1.add_run("Email: ").bold = True
p1.add_run("priya.nair@techcorp.in\n")
p1.add_run("Phone: ").bold = True
p1.add_run("+91-8041239870\n")
p1.add_run("Category: ").bold = True
p1.add_run("Enterprise Fiber Dedicated Lease Line (Order ID: FL-4401)")

doc2.add_heading("Grievance Description", level=2)
doc2.add_paragraph(
    "Our development hub in Whitefield has experienced a complete commercial internet outage since 09:00 AM IST today. "
    "Our contract SLA tier guarantees 99.9% network availability with a mandatory 2-hour response window for critical priority outages. "
    "We have attached network monitoring logs, traceroute outputs, and router error dumps for your tier-3 engineering group."
)

doc2.add_heading("Status & Action Taken", level=2)
doc2.add_paragraph(
    "Field engineer assigned ticket #ENG-4412. The splicing team is actively on-site inspecting physical fiber cuts near the substation. "
    "Emergency cellular backup unit was dispatched. The case is currently In Progress."
)

file2_path = DATA_DIR / "complaint_002.docx"
doc2.save(str(file2_path))
print(f"Generated: {file2_path}")


# -------------------------------------------------------------
# Document 3: PDF File (.pdf) - Defective Hardware Delivery
# -------------------------------------------------------------
class CleanPDF(FPDF):
    def header(self):
        self.set_font("Helvetica", "B", 14)
        self.cell(0, 10, "CUSTOMER GRIEVANCE & DAMAGE NOTICE", border=False, new_x="LMARGIN", new_y="NEXT", align="C")
        self.ln(5)

pdf = CleanPDF()
pdf.add_page()
pdf.set_font("Helvetica", size=10)

pdf.set_font("Helvetica", "B", 11)
pdf.cell(0, 6, "Order Number: ORD-1092834", new_x="LMARGIN", new_y="NEXT")
pdf.cell(0, 6, "Customer Name: Amit Verma", new_x="LMARGIN", new_y="NEXT")
pdf.cell(0, 6, "Email: amit.verma@freemail.com", new_x="LMARGIN", new_y="NEXT")
pdf.cell(0, 6, "Phone: +91-9823011223", new_x="LMARGIN", new_y="NEXT")
pdf.cell(0, 6, "Category: Hardware Delivery", new_x="LMARGIN", new_y="NEXT")
pdf.ln(5)

pdf.set_font("Helvetica", "B", 11)
pdf.cell(0, 6, "Description of Issue:", new_x="LMARGIN", new_y="NEXT")
pdf.set_font("Helvetica", size=10)
pdf.multi_cell(0, 6, (
    "I ordered an ergonomic dual-monitor desktop arm on August 1st. The delivery arrived yesterday via express courier, "
    "but upon unboxing, the main mounting bracket was severely bent and the packet containing M4 mounting screws was missing. "
    "I am currently unable to mount either display.\n\n"
    "I have included photographs of the damaged shipping carton and the invoice receipt. "
    "I do not want an escalation or dispute; I simply request a replacement bracket or the missing accessory pack dispatched."
))
pdf.ln(5)

pdf.set_font("Helvetica", "B", 11)
pdf.cell(0, 6, "Resolution & Case Status:", new_x="LMARGIN", new_y="NEXT")
pdf.set_font("Helvetica", size=10)
pdf.multi_cell(0, 6, (
    "Customer service rep acknowledged hardware transit damage. A complimentary replacement unit (AWB #8831920) "
    "was dispatched this morning via courier. Case marked Resolved upon tracking link transmission."
))

file3_path = DATA_DIR / "complaint_003.pdf"
pdf.output(str(file3_path))
print(f"Generated: {file3_path}")


# -------------------------------------------------------------
# Document 4: Word Document (.docx) - Account Security & Login Failure
# -------------------------------------------------------------
doc4 = docx.Document()
doc4.add_heading("SECURITY INCIDENT / LOCKED ACCOUNT ESCALATION", level=1)

p4 = doc4.add_paragraph()
p4.add_run("Customer Name: ").bold = True
p4.add_run("Kavita Sundaram\n")
p4.add_run("Email: ").bold = True
p4.add_run("kavita.sundaram@finnet.org\n")
p4.add_run("Phone: ").bold = True
p4.add_run("+91-9123456789\n")
p4.add_run("Category: ").bold = True
p4.add_run("User Account & Security Access")

doc4.add_heading("Details of Incident", level=2)
doc4.add_paragraph(
    "My corporate administrator account was locked out after multiple unauthorized MFA prompts were received on my registered mobile device. "
    "Our payroll and internal accounting tasks are completely blocked. We suspect a credential stuffing attempt. "
    "I have attached our audit access log showing IP anomalies from external regions."
)

doc4.add_heading("Action Taken & Severity", level=2)
doc4.add_paragraph(
    "Security operations revoked active API tokens and initiated an account verification review. "
    "Requires immediate senior administrator escalation to re-verify credentials and restore single sign-on access. "
    "Status: Open."
)

file4_path = DATA_DIR / "complaint_004.docx"
doc4.save(str(file4_path))
print(f"Generated: {file4_path}")


# -------------------------------------------------------------
# Document 5: PDF File (.pdf) - Minor Query / Non-Complaint
# -------------------------------------------------------------
pdf5 = CleanPDF()
pdf5.add_page()
pdf5.set_font("Helvetica", size=10)

pdf5.set_font("Helvetica", "B", 11)
pdf5.cell(0, 6, "Account Inquiry: INQ-33102", new_x="LMARGIN", new_y="NEXT")
pdf5.cell(0, 6, "Customer Name: Vikram Malhotra", new_x="LMARGIN", new_y="NEXT")
pdf5.cell(0, 6, "Email: vikram.m@domain.co", new_x="LMARGIN", new_y="NEXT")
pdf5.cell(0, 6, "Category: General Inquiry", new_x="LMARGIN", new_y="NEXT")
pdf5.ln(5)

pdf5.set_font("Helvetica", "B", 11)
pdf5.cell(0, 6, "Inquiry Details:", new_x="LMARGIN", new_y="NEXT")
pdf5.set_font("Helvetica", size=10)
# Replaced em-dash (\u2014) with standard ASCII hyphen (-)
pdf5.multi_cell(0, 6, (
    "Hello team, I am writing to inquire whether our enterprise tier subscription includes the new batch API features "
    "announced in the recent developer changelog. This is not a grievance or complaint - we are simply planning our Q4 infrastructure "
    "budget and would like documentation regarding rate limits for bulk webhook processing.\n\n"
    "No supporting documents attached as none are needed."
))
pdf5.ln(5)

pdf5.set_font("Helvetica", "B", 11)
pdf5.cell(0, 6, "Current Status:", new_x="LMARGIN", new_y="NEXT")
pdf5.set_font("Helvetica", size=10)
pdf5.multi_cell(0, 6, (
    "Documentation link sent to customer email. Ticket closed as resolved inquiry."
))

file5_path = DATA_DIR / "complaint_005.pdf"
pdf5.output(str(file5_path))
print(f"Generated: {file5_path}")